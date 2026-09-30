"""Import CSV/Excel avec validation, détection doublons, rapport d'erreurs — offline."""
import csv
import os

EXPECTED = ["nom", "reference", "code_produit", "categorie", "marque",
            "prix_achat", "prix_vente", "prix_promo", "quantite", "unite",
            "description", "code_barres", "format_barres"]

ALIASES = {
    "name": "nom", "product": "nom", "produit": "nom", "designation": "nom",
    "ref": "reference", "référence": "reference",
    "price": "prix_vente", "prix": "prix_vente", "prix vente": "prix_vente",
    "buy": "prix_achat", "prix achat": "prix_achat",
    "promo": "prix_promo", "qty": "quantite", "qte": "quantite",
    "unit": "unite", "barcode": "code_barres", "code-barres": "code_barres",
    "codebarre": "code_barres", "category": "categorie", "brand": "marque",
}


def _norm_header(h):
    h = (h or "").strip().lower()
    return ALIASES.get(h, h)


def read_file(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".csv":
        # détection délimiteur + encodage
        for enc in ("utf-8-sig", "utf-8", "cp1252", "latin1"):
            try:
                with open(path, "r", encoding=enc, newline="") as f:
                    sample = f.read(4096)
                    f.seek(0)
                    try:
                        dialect = csv.Sniffer().sniff(sample, delimiters=";,,\t|")
                        delim = dialect.delimiter
                    except Exception:
                        delim = ";" if sample.count(";") > sample.count(",") else ","
                    reader = csv.DictReader(f, delimiter=delim)
                    rows = list(reader)
                    headers = [_norm_header(h) for h in (reader.fieldnames or [])]
                    normed = []
                    for r in rows:
                        normed.append({_norm_header(k): (v or "").strip() if isinstance(v, str) else v
                                       for k, v in r.items()})
                    return normed, headers
            except UnicodeError:
                continue
        raise ValueError("Encodage CSV illisible.")
    elif ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return [], []
        headers = [_norm_header(str(h or "")) for h in rows[0]]
        data = []
        for r in rows[1:]:
            d = {}
            for h, v in zip(headers, r):
                d[h] = "" if v is None else str(v).strip()
            data.append(d)
        return data, headers
    elif ext == ".xls":
        raise ValueError("Format .xls ancien non supporté — convertissez en .xlsx ou .csv.")
    raise ValueError("Format non supporté (utilisez .csv ou .xlsx).")


def validate_rows(rows, db):
    """Retourne (valides, erreurs). erreurs: liste {ligne, données, erreur}."""
    valides, erreurs = [], []
    vus_cb, vus_ref = set(), set()
    for i, r in enumerate(rows, start=2):
        try:
            nom = (r.get("nom") or "").strip()
            if not nom:
                raise ValueError("Nom manquant.")
            pv_raw = (r.get("prix_vente") or "0").strip().replace(" ", "").replace(",", ".")
            pa_raw = (r.get("prix_achat") or "0").strip().replace(" ", "").replace(",", ".")
            try:
                pv = float(pv_raw or 0)
                pa = float(pa_raw or 0)
            except Exception:
                raise ValueError(f"Prix invalide (achat='{pa_raw}', vente='{pv_raw}').")
            if pv < 0 or pa < 0:
                raise ValueError("Prix négatif.")
            pp = (r.get("prix_promo") or "").strip().replace(" ", "").replace(",", ".")
            pp_f = None
            if pp:
                try:
                    pp_f = float(pp)
                except Exception:
                    raise ValueError(f"Prix promo invalide: '{pp}'.")
                if pp_f < 0:
                    raise ValueError("Prix promo négatif.")
            q_raw = (r.get("quantite") or "0").strip().replace(" ", "").replace(",", ".")
            try:
                q = float(q_raw or 0)
            except Exception:
                raise ValueError(f"Quantité invalide: '{q_raw}'.")
            cb = (r.get("code_barres") or "").strip()
            ref = (r.get("reference") or "").strip()
            if cb:
                if cb in vus_cb:
                    raise ValueError(f"Doublon dans le fichier (code-barres {cb}).")
                if db.find_by_barcode(cb):
                    raise ValueError(f"Code-barres déjà en base: {cb}.")
                vus_cb.add(cb)
            if ref:
                if ref in vus_ref:
                    raise ValueError(f"Doublon dans le fichier (référence {ref}).")
                vus_ref.add(ref)
            valides.append({
                "nom": nom, "reference": ref or None, "code_produit": (r.get("code_produit") or "").strip(),
                "categorie": (r.get("categorie") or "").strip(), "marque": (r.get("marque") or "").strip(),
                "prix_achat": pa, "prix_vente": pv, "prix_promo": pp_f,
                "quantite": q, "unite": (r.get("unite") or "pcs").strip(),
                "description": (r.get("description") or "").strip(),
                "code_barres": cb or None,
                "format_barres": (r.get("format_barres") or "EAN13").strip().upper() or "EAN13",
            })
        except Exception as e:
            erreurs.append({"ligne": i, "donnees": r, "erreur": str(e)})
    return valides, erreurs


def import_validated(db, valides):
    ok, ko = 0, []
    for v in valides:
        try:
            db.add_product(v)
            ok += 1
        except Exception as e:
            ko.append({"donnees": v, "erreur": str(e)})
    return ok, ko
